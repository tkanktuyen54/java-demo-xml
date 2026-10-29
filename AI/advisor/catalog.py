"""Strict Excel/CSV import. Importing never changes purchase rights or website prices."""

import argparse
import csv
import json
import os
import re
import tempfile
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from .schemas import Course


def normalize(text: str) -> str:
  text = unicodedata.normalize("NFD", str(text).lower().replace("đ", "d"))
  return "".join(char for char in text if unicodedata.category(char) != "Mn")


HEADERS = {
  "ma khoa hoc": "slug",
  "ten khoa hoc": "title",
  "gia": "price",
  "gia vnd": "price",
  "trinh do": "level",
  "danh muc": "category",
  "mo ta": "description",
  "giang vien": "instructor",
  "thoi luong": "hours",
  "ket qua dat duoc": "outcomes",
  "doi tuong": "audience",
  "yeu cau dau vao": "prerequisites",
}


def parse_price(value) -> Decimal:
  if isinstance(value, bool) or value is None:
    raise ValueError("Giá phải là số tiền VND, ví dụ 499000.")
  if isinstance(value, (int, float, Decimal)):
    price = Decimal(str(value))
  else:
    text = str(value).strip().upper()
    text = re.sub(r"\s*(VND|VNĐ|Đ|₫)\s*$", "", text).strip()
    if re.fullmatch(r"\d{1,3}(?:[.,]\d{3})+", text):
      text = text.replace(".", "").replace(",", "")
    if not re.fullmatch(r"\d+", text):
      raise ValueError("Giá phải là số nguyên VND hoặc dạng 499.000 ₫.")
    try:
      price = Decimal(text)
    except InvalidOperation as ex:
      raise ValueError("Giá không hợp lệ.") from ex
  if not price.is_finite() or price < 0 or price != price.to_integral_value():
    raise ValueError("Giá phải là số nguyên VND không âm.")
  return price


def read_spreadsheet(path: Path, sheet: str = None) -> list:
  if path.suffix.lower() == ".csv":
    with path.open(encoding="utf-8-sig", newline="") as source:
      rows = list(csv.reader(source))
  elif path.suffix.lower() == ".xlsx":
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
      worksheet = workbook[sheet] if sheet else workbook.active
      rows = list(worksheet.iter_rows(values_only=True))
    finally:
      workbook.close()
  else:
    raise ValueError("Chỉ nhận .xlsx hoặc .csv UTF-8.")
  if not rows:
    raise ValueError("File không có dữ liệu.")
  columns = [HEADERS.get(normalize(cell).strip(), str(cell or "").strip()) for cell in rows[0]]
  if len(set(columns)) != len(columns):
    raise ValueError("File có tiêu đề cột trùng nhau.")
  if not {"slug", "title", "price"}.issubset(columns):
    raise ValueError("Thiếu cột slug, title, price (hoặc Mã khóa học, Tên khóa học, Giá).")
  courses, seen = [], set()
  for row_number, row in enumerate(rows[1:], start=2):
    if all(value is None or value == "" for value in row):
      continue
    if len(courses) >= 5000:
      raise ValueError("Mỗi file tối đa 5.000 khóa học.")
    values = {key: value for key, value in zip(columns, row) if value is not None and value != ""}
    try:
      if any(isinstance(value, str) and value.startswith("=") for value in values.values()):
        raise ValueError("Hãy thay công thức bằng giá trị đã xác nhận trước khi nhập.")
      values["price"] = parse_price(values.get("price"))
      if isinstance(values.get("outcomes"), str):
        values["outcomes"] = [
          value.strip() for value in values["outcomes"].split("|") if value.strip()
        ]
      course = Course.model_validate(values)
      if course.slug in seen:
        raise ValueError(f"Trùng mã khóa học: {course.slug}")
      seen.add(course.slug)
      courses.append(course)
    except (ValueError, TypeError) as ex:
      raise ValueError(f"Dòng {row_number}: {ex}") from ex
  if not courses:
    raise ValueError("File không có khóa học hợp lệ.")
  return courses


def load_catalog(path: Path) -> list:
  return [Course.model_validate(row) for row in json.loads(path.read_text(encoding="utf-8"))]


def main():
  parser = argparse.ArgumentParser(description="Nhập Excel/CSV thành danh mục tư vấn JSON.")
  parser.add_argument("file", type=Path)
  parser.add_argument("--sheet")
  parser.add_argument("--output", type=Path, default=Path("data/catalog.json"))
  args = parser.parse_args()
  try:
    courses = read_spreadsheet(args.file, args.sheet)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
      mode="w", encoding="utf-8", dir=args.output.parent, delete=False
    ) as target:
      temporary = target.name
      json.dump(
        [course.model_dump(mode="json") for course in courses], target, ensure_ascii=False, indent=2
      )
      target.write("\n")
    os.replace(temporary, args.output)
  except (ValueError, OSError, KeyError) as ex:
    parser.exit(1, f"Nhập thất bại: {ex}\n")
  print(f"Đã nhập {len(courses)} khóa học vào {args.output}. Khởi động lại AI để nạp dữ liệu mới.")


if __name__ == "__main__":
  main()
